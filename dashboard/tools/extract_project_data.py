"""Extract the Team 1 (DL) project plan from PeakDeconv_Gantt_Chart.xlsx.

Reads ONLY the "Task overview" and "Team 1 · DL" sheets, cross-checks them against
each other and writes ../data/project-data.js (a plain JS module that sets
window.PD_DATA). The combined "Gantt chart" sheet and the other team's sheet are
never opened.

Run from the dashboard folder:
    python tools/extract_project_data.py
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "data" / "source" / "PeakDeconv_Gantt_Chart.xlsx"
OUT = ROOT / "data" / "project-data.js"

WEEK_COL0 = 3          # column C = week 1
N_WEEKS = 15
OWNER_KEYS = {"BME lead": "bme", "EEE lead": "eee", "BME+EEE": "shared"}
THICK = {"thick", "medium"}


def parse_weeks(text):
    """'Weeks 5–7, Weeks 9–12' -> [5,6,7,9,10,11,12]; 'Week 2' -> [2]."""
    weeks = []
    for part in text.split(","):
        nums = [int(n) for n in re.findall(r"\d+", part)]
        if len(nums) == 1:
            weeks.append(nums[0])
        elif len(nums) == 2:
            weeks.extend(range(nums[0], nums[1] + 1))
        else:
            raise ValueError(f"cannot parse weeks: {text!r}")
    return weeks


def segments(weeks):
    """[5,6,7,9,10] -> [[5,7],[9,10]]"""
    segs = []
    for w in sorted(weeks):
        if segs and w == segs[-1][1] + 1:
            segs[-1][1] = w
        else:
            segs.append([w, w])
    return segs


def is_thick(cell):
    b = cell.border
    sides = [b.left.style, b.right.style, b.top.style, b.bottom.style]
    return sum(1 for s in sides if s in THICK) >= 3


def main():
    wb = openpyxl.load_workbook(XLSX)
    ov = wb["Task overview"]
    t1 = wb["Team 1 · DL"]

    # ---- week header (row 8 numbers, row 9 labels) ------------------------------
    weeks = []
    for i in range(N_WEEKS):
        col = WEEK_COL0 + i
        num = t1.cell(8, col).value
        label = str(t1.cell(9, col).value).strip().upper()
        assert num == i + 1, f"week header mismatch at col {col}"
        kind = {"WEEK": "work", "PRESENTATION": "presentation", "MIDTERM": "midterm"}[label]
        weeks.append({"week": num, "kind": kind})

    # ---- Task overview: normalized task list ------------------------------------
    packages, tasks = [], []
    current = None
    for r in range(5, ov.max_row + 1):
        a, b, c, d, e = (ov.cell(r, k).value for k in range(1, 6))
        if a and b is None and " · " in str(a):
            code, title = str(a).split(" · ", 1)
            current = {"id": code.strip(), "title": title.strip()}
            packages.append(current)
            continue
        if a and b:
            owner_line = str(e).split("\n")[0].strip()
            tasks.append({
                "id": str(a).strip(),
                "wp": current["id"],
                "name": str(b).strip(),
                "weeksLabel": str(c).strip(),
                "weeks": parse_weeks(str(c)),
                "duration": int(d),
                "ownerLabel": owner_line,
                "owner": OWNER_KEYS[owner_line],
            })

    # ---- Team 1 · DL grid: cross-check + DL rows, fixed delivery, WP finals -------
    by_id = {t["id"]: t for t in tasks}
    wp_meta = {}
    seen_grid = set()
    row = 10
    while row <= 44:
        code = t1.cell(row, 1).value
        name = t1.cell(row, 2).value
        if code and re.fullmatch(r"WP\d", str(code)):
            finals = [WEEK_COL0 + i for i in range(N_WEEKS)
                      if t1.cell(row, WEEK_COL0 + i).value == "◆"]
            assert len(finals) == 1, f"{code}: expected one ◆"
            bar = [i + 1 for i in range(N_WEEKS)
                   if t1.cell(row, WEEK_COL0 + i).fill.fill_type == "solid"
                   and t1.cell(row, WEEK_COL0 + i).fill.fgColor.rgb not in ("FF1F2B38",)]
            accent = t1.cell(row, 1).border.left.color.rgb[2:]
            ink = t1.cell(row, 1).font.color.rgb[2:]
            tint = t1.cell(row, 1).fill.fgColor.rgb[2:]
            wp_meta[str(code)] = {"name": str(name).strip(), "finalWeek": finals[0] - WEEK_COL0 + 1,
                                  "headingBar": bar, "xlsxAccent": accent, "xlsxInk": ink, "xlsxTint": tint}
            row += 1
            continue
        if code and str(code) in by_id:
            t = by_id[str(code)]
            assert str(name).strip() == t["name"], f"name mismatch {code}"
            cells = [(i + 1, t1.cell(row, WEEK_COL0 + i)) for i in range(N_WEEKS)]
            active = [w for w, c in cells if c.value]
            assert active == t["weeks"], f"{code}: grid {active} != overview {t['weeks']}"
            assert len(active) == t["duration"], f"{code}: duration mismatch"
            owners = {str(c.value).strip() for _, c in cells if c.value}
            assert owners == {t["ownerLabel"]}, f"{code}: owner mismatch {owners}"
            fixed = [w for w, c in cells if c.value and is_thick(c)]
            assert len(fixed) == 1, f"{code}: expected one fixed delivery cell, got {fixed}"
            t["fixedWeek"] = fixed[0]
            t["segments"] = segments(active)
            # helper row directly below: "Team 1 · DL" metadata, not a task
            nxt = row + 1
            helper = [t1.cell(nxt, WEEK_COL0 + i).value for i in range(N_WEEKS)]
            if t1.cell(nxt, 1).value is None and any(helper):
                dl_weeks = [i + 1 for i, v in enumerate(helper) if v]
                assert all(str(v).strip() == "Team 1 · DL" for v in helper if v)
                assert dl_weeks == active, f"{code}: DL row weeks differ"
                t["dl"] = True
                row += 2
            else:
                t["dl"] = False
                row += 1
            seen_grid.add(t["id"])
            continue
        raise AssertionError(f"unexpected row {row}: {code!r} {name!r}")

    assert seen_grid == set(by_id), "grid and overview task sets differ"
    for p in packages:
        m = wp_meta[p["id"]]
        assert m["name"] == p["title"], f"{p['id']} title mismatch"
        union = sorted({w for t in tasks if t["wp"] == p["id"] for w in t["weeks"]})
        if union != m["headingBar"]:
            print(f"WARNING {p['id']}: workbook heading bar covers weeks {m['headingBar']}, "
                  f"subtasks cover {union} — both kept (headingBar / taskWeeks)")
        p.update(finalWeek=m["finalWeek"], headingBar=m["headingBar"], taskWeeks=union,
                 xlsxAccent=m["xlsxAccent"], xlsxInk=m["xlsxInk"], xlsxTint=m["xlsxTint"])

    # legend short names ("■ WP1 · Acquisition")
    for r in range(51, 54):
        for col in range(1, 18):
            v = t1.cell(r, col).value
            if v and "WP" in str(v):
                code, short = str(v).replace("■", "").strip().split(" · ")
                next(p for p in packages if p["id"] == code)["short"] = short

    # ---- team (Team 1 sheet only) -----------------------------------------------
    assert t1["D47"].value == "Team 1"
    team = [{"department": t1["C48"].value, "name": t1["D48"].value},
            {"department": t1["C49"].value, "name": t1["D49"].value}]

    subtitle = t1["O2"].value  # "15 weeks · midterm in week 8"

    data = {
        "project": {
            "title": t1["A2"].value,
            "subtitle": t1["A3"].value.split(" · ")[0],
            "scheduleNote": subtitle,
            "track": {"team": "Team 1", "approach": "Deep Learning", "short": "DL"},
        },
        "weeks": weeks,
        "packages": packages,
        "tasks": tasks,
        "team": team,
        "owners": {
            "bme": {"label": "BME lead", "meaning": "BME primary, EEE support"},
            "eee": {"label": "EEE lead", "meaning": "EEE primary, BME support"},
            "shared": {"label": "BME + EEE", "meaning": "Shared"},
        },
        "notes": {
            "finalWeek": "◆ marks the planned final week of a work package.",
            "fixedWeek": "Fixed delivery week — the earlier weeks of a task may be rearranged, not this one.",
            "chatbot": "WP6.3 chatbot: single task, result queries only.",
        },
        "source": {"file": XLSX.name, "sheets": ["Task overview", "Team 1 · DL"]},
    }

    text = json.dumps(data, ensure_ascii=False, indent=2)
    assert not re.search(r"Team\s*2|\bML\b", text), "other-team content leaked into data"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(
        "// Generated by tools/extract_project_data.py from data/source/PeakDeconv_Gantt_Chart.xlsx — do not edit by hand.\n"
        "window.PD_DATA = " + text + ";\n", encoding="utf-8")

    n_dl = sum(t["dl"] for t in tasks)
    print(f"weeks={len(weeks)} packages={len(packages)} tasks={len(tasks)} dl={n_dl} "
          f"task-weeks={sum(t['duration'] for t in tasks)} -> {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
